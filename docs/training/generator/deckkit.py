# -*- coding: utf-8 -*-
"""Helpers for building the manual-workflow training deck with python-pptx."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- palette
INK        = "22303C"
BODY       = "44515C"
MUTED      = "8494A0"
ACCENT     = "D2542F"
ACCENT_DK  = "9B3A1D"
ACCENT_TNT = "FAEAE3"
PANEL      = "F1F4F7"
PANEL_DK   = "E3E9ED"
LINE       = "D3DBE1"
GREEN      = "2F6B4F"
GREEN_TNT  = "E7F1EB"
CODE_BG    = "1B262E"
CODE_FG    = "DCE6EC"
CODE_KEY   = "EFA277"
WHITE      = "FFFFFF"
DARK_FG    = "E9EEF2"

KO   = "Malgun Gothic"
MONO = "Consolas"

SW, SH = 13.333, 7.5
ML, MR = 0.72, 0.72
CW = SW - ML - MR        # content width 11.893


def new_deck():
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    return prs


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def rgb(h):
    return RGBColor.from_string(h)


# ---------------------------------------------------------------- xml bits
def _set_run_font(run, name):
    rPr = run._r.get_or_add_rPr()
    for tag, attr in (("a:latin", "typeface"), ("a:ea", "typeface"),
                      ("a:cs", "typeface")):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set(attr, name)


def _bullet(p, char="▪", color=ACCENT, indent=0.22):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(Emu(int(Inches(indent)))))
    pPr.set("indent", str(-Emu(int(Inches(indent)))))
    clr = etree.SubElement(pPr, qn("a:buClr"))
    srgb = etree.SubElement(clr, qn("a:srgbClr"))
    srgb.set("val", color)
    fnt = etree.SubElement(pPr, qn("a:buFont"))
    fnt.set("typeface", "Arial")
    ch = etree.SubElement(pPr, qn("a:buChar"))
    ch.set("char", char)


def _no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    etree.SubElement(pPr, qn("a:buNone"))


# ---------------------------------------------------------------- text
def textbox(slide, x, y, w, h, anchor="t", wrap=True):
    # 음수 크기는 PowerPoint가 파일 자체를 거부하므로 항상 양수로 막습니다.
    w, h = max(w, 0.1), max(h, 0.1)
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE,
                          "b": MSO_ANCHOR.BOTTOM}[anchor]
    return tb, tf


def para(tf, first=False):
    return tf.paragraphs[0] if first else tf.add_paragraph()


def style(p, size=14, color=BODY, bold=False, font=KO, align="l",
          space_after=0, space_before=0, line=1.28):
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                   "r": PP_ALIGN.RIGHT}[align]
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    for r in p.runs:
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = rgb(color)
        _set_run_font(r, font)
    return p


def style_para(p, align="l", space_after=0, space_before=0, line=1.28):
    """런은 이미 개별로 꾸며진 상태이므로 단락 속성만 건드립니다."""
    p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER,
                   "r": PP_ALIGN.RIGHT}[align]
    p.space_after = Pt(space_after)
    p.space_before = Pt(space_before)
    p.line_spacing = line
    return p


def write(slide, x, y, w, h, spans, size=14, color=BODY, bold=False,
          font=KO, align="l", line=1.28, anchor="t"):
    """spans: str | list of paragraph specs.

    Each paragraph spec: str, or dict with
      t: str or list[(text, dict-overrides)], and per-paragraph overrides.
    """
    _, tf = textbox(slide, x, y, w, h, anchor=anchor)
    if isinstance(spans, str):
        spans = [spans]
    for i, spec in enumerate(spans):
        p = para(tf, first=(i == 0))
        if isinstance(spec, str):
            spec = {"t": spec}
        opts = dict(size=size, color=color, bold=bold, font=font,
                    align=align, line=line,
                    space_after=spec.get("sa", 0),
                    space_before=spec.get("sb", 0))
        for k in ("size", "color", "bold", "font", "align", "line"):
            if k in spec:
                opts[k] = spec[k]
        chunks = spec["t"]
        if isinstance(chunks, str):
            chunks = [(chunks, {})]
        for text, over in chunks:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(over.get("size", opts["size"]))
            r.font.bold = over.get("bold", opts["bold"])
            r.font.italic = over.get("italic", False)
            r.font.color.rgb = rgb(over.get("color", opts["color"]))
            _set_run_font(r, over.get("font", opts["font"]))
        style_para(p, align=opts["align"], line=opts["line"],
                   space_after=opts["space_after"],
                   space_before=opts["space_before"])
        if spec.get("bullet"):
            _bullet(p, char=spec.get("bullet_char", "▪"),
                    color=spec.get("bullet_color", ACCENT),
                    indent=spec.get("indent", 0.22))
        else:
            _no_bullet(p)
    return tf


def bullets(slide, x, y, w, h, items, size=13.5, color=BODY, gap=9,
            char="▪", bullet_color=ACCENT, line=1.3):
    specs = []
    for it in items:
        spec = {"bullet": True, "bullet_char": char,
                "bullet_color": bullet_color, "sa": gap}
        spec["t"] = it if not isinstance(it, dict) else it["t"]
        if isinstance(it, dict):
            for k in ("size", "color", "bold", "sa", "bullet",
                      "bullet_char", "bullet_color", "indent"):
                if k in it:
                    spec[k] = it[k]
        specs.append(spec)
    return write(slide, x, y, w, h, specs, size=size, color=color, line=line)


# ---------------------------------------------------------------- shapes
def rect(slide, x, y, w, h, fill=PANEL, line_color=None, radius=None,
         shadow=False, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    w, h = max(w, 0.1), max(h, 0.1)
    sh = slide.shapes.add_shape(shape, Inches(x), Inches(y),
                                Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
    if line_color is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = rgb(line_color)
        sh.line.width = Pt(1)
    if not shadow:
        sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.text_frame.word_wrap = True
    return sh


def circle(slide, cx, cy, d, fill=ACCENT):
    return rect(slide, cx - d / 2, cy - d / 2, d, d, fill=fill,
                shape=MSO_SHAPE.OVAL)


def badge(slide, x, y, d, label, fill=ACCENT, color=WHITE, size=17,
          shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.28, font=KO):
    sh = rect(slide, x, y, d, d, fill=fill, radius=radius, shape=shape)
    tf = sh.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = str(label)
    style(p, size=size, color=color, bold=True, align="c", font=font, line=1.0)
    _no_bullet(p)
    return sh


def arrow(slide, x, y, w, h, color=MUTED, shape=MSO_SHAPE.RIGHT_ARROW):
    sh = rect(slide, x, y, w, h, fill=color, shape=shape)
    return sh


def code(slide, x, y, w, lines, size=12, pad=0.2, lh=0.235, bg=CODE_BG,
         fg=CODE_FG):
    h = pad * 2 + lh * len(lines)
    box = rect(slide, x, y, w, h, fill=bg, radius=0.06)
    _, tf = textbox(slide, x + 0.22, y + pad - 0.03, w - 0.44,
                    h - pad * 2 + 0.1)
    for i, ln in enumerate(lines):
        p = para(tf, first=(i == 0))
        chunks = ln if isinstance(ln, list) else [(ln, {})]
        for text, over in chunks:
            r = p.add_run()
            r.text = text
            r.font.size = Pt(over.get("size", size))
            r.font.bold = over.get("bold", False)
            r.font.color.rgb = rgb(over.get("color", fg))
            _set_run_font(r, over.get("font", MONO))
        style_para(p, line=1.15, space_after=0)
        _no_bullet(p)
    return box, h


# ---------------------------------------------------------------- table
NO_STYLE_GRID = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # No Style, No Grid


def table(slide, x, y, w, col_w, rows, head_h=0.42, row_h=0.36,
          size=11.5, head_size=11.5, head_fill=INK, head_color=WHITE,
          zebra=(WHITE, PANEL), body_color=BODY, aligns=None,
          bold_first_col=False):
    nrows, ncols = len(rows), len(rows[0])
    h = head_h + row_h * (nrows - 1)
    gfx = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y),
                                 Inches(w), Inches(h))
    tbl = gfx.table
    # kill the default banded blue style
    tblPr = tbl._tbl.find(qn("a:tblPr"))
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = NO_STYLE_GRID

    total = sum(col_w)
    for j, cw in enumerate(col_w):
        tbl.columns[j].width = Emu(int(Inches(w) * cw / total))
    tbl.rows[0].height = Inches(head_h)
    for i in range(1, nrows):
        tbl.rows[i].height = Inches(row_h)

    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = tbl.cell(i, j)
            cell.margin_left = Inches(0.11)
            cell.margin_right = Inches(0.09)
            cell.margin_top = Inches(0.045)
            cell.margin_bottom = Inches(0.045)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i == 0:
                cell.fill.fore_color.rgb = rgb(head_fill)
            else:
                cell.fill.fore_color.rgb = rgb(zebra[(i - 1) % 2])
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            chunks = val if isinstance(val, list) else [(str(val), {})]
            for text, over in chunks:
                r = p.add_run()
                r.text = text
                r.font.size = Pt(over.get(
                    "size", head_size if i == 0 else size))
                r.font.bold = over.get(
                    "bold", True if i == 0 or (bold_first_col and j == 0)
                    else False)
                r.font.color.rgb = rgb(over.get(
                    "color", head_color if i == 0 else body_color))
                _set_run_font(r, over.get("font", KO))
            al = "l"
            if aligns:
                al = aligns[j]
            style_para(p, align=al, line=1.2)
            _no_bullet(p)
    return gfx, h


# ---------------------------------------------------------------- chrome
STATE = {"n": 0, "total": 0}


def title_block(slide, title, lead=None, step=None, y=0.5):
    x = ML
    if step is not None:
        badge(slide, ML, y - 0.02, 0.66, step, size=20)
        x = ML + 0.92
    write(slide, x, y - 0.06, CW - (x - ML), 0.62, title, size=31,
          color=INK, bold=True, line=1.0)
    if lead:
        write(slide, x, y + 0.56, CW - (x - ML), 0.5, lead, size=13.5,
              color=MUTED, line=1.25)
    return y + (1.16 if lead else 0.8)


def footer(slide, label=""):
    STATE["n"] += 1
    if label:
        write(slide, ML, SH - 0.52, CW - 1.2, 0.26, label, size=9.5,
              color=MUTED)
    write(slide, SW - MR - 1.2, SH - 0.52, 1.2, 0.26,
          "%d" % STATE["n"], size=9.5, color=MUTED, align="r")


def dark_bg(slide, color=INK):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0,
                                Inches(SW), Inches(SH))
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(color)
    sh.line.fill.background()
    sh.shadow.inherit = False
    return sh


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text
